"""
Importación masiva de clientes desde una plantilla de Excel.

Este módulo concentra dos responsabilidades: generar la plantilla que el
usuario descarga y validar/importar el archivo diligenciado. Cada fila se
valida antes de guardarse (RQF013); las filas con error se informan con su
número y motivo, y las válidas se importan aunque otras fallen.
"""

from io import BytesIO

from django.contrib.auth import get_user_model
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Cliente, ClienteMunicipio, ClienteServicio, TipoServicio

Usuario = get_user_model()

# Relación entre la clave interna de cada dato y el encabezado de la plantilla.
COLUMNAS = [
    ("nombre", "Nombre o razón social", 32),
    ("nit", "Número de documento", 20),
    ("tipo_documento", "Tipo de documento", 22),
    ("tipo_persona", "Tipo de persona", 16),
    ("regimen", "Régimen tributario", 24),
    ("responsable_iva", "Responsable de IVA", 20),
    ("ciudad", "Ciudad", 20),
    ("municipio_principal", "Municipio principal", 22),
    ("telefono", "Teléfono del cliente", 22),
    ("correo_cliente", "Correo del cliente", 30),
    ("responsable_correo", "Correo del contador (asignado en Contanika)", 36),
    ("tipos_servicio", "Tipos de servicio", 30),
    ("municipios_adicionales", "Municipios adicionales", 30),
]

EJEMPLO = {
    "nombre": "Comercializadora La Rebaja SAS",
    "nit": "900123456",
    "tipo_documento": "NIT",
    "tipo_persona": "Jurídica",
    "regimen": "Régimen ordinario",
    "responsable_iva": "Sí",
    "ciudad": "Bogotá D.C.",
    "municipio_principal": "Bogotá D.C.",
    "telefono": "3101234567",
    "correo_cliente": "contacto@larebaja.co",
    "responsable_correo": "",
    "tipos_servicio": "Declaración de renta; Declaración de IVA",
    "municipios_adicionales": "Soacha",
}

VALORES_VERDADEROS = {"si", "sí", "s", "1", "true", "verdadero", "x"}
VALORES_FALSOS = {"no", "n", "0", "false", "falso", ""}


def generar_plantilla_clientes():
    """Construye la plantilla de Excel con encabezados, ejemplo e instrucciones."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Clientes"

    encabezado = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="15616D")
    for indice, (_clave, titulo, ancho) in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(row=1, column=indice, value=titulo)
        celda.font = encabezado
        celda.fill = relleno
        celda.alignment = Alignment(horizontal="center")
        hoja.column_dimensions[get_column_letter(indice)].width = ancho

    for indice, (_clave, _titulo, _ancho) in enumerate(COLUMNAS, start=1):
        clave = COLUMNAS[indice - 1][0]
        hoja.cell(row=2, column=indice, value=EJEMPLO[clave])

    instrucciones = libro.create_sheet("Instrucciones")
    instrucciones.column_dimensions["A"].width = 26
    instrucciones.column_dimensions["B"].width = 70
    textos = [
        ("Columna", "Qué debe llevar"),
        ("Nombre o razón social", "Obligatorio. Nombre completo o razón social del cliente."),
        ("NIT o cédula", "Obligatorio. Sin puntos ni espacios. No puede repetirse."),
        ("Tipo de persona", "Obligatorio. Escriba Natural o Jurídica."),
        ("Régimen tributario", "Régimen del cliente (se puede escribir uno nuevo)."),
        ("Responsable de IVA", "Escriba Sí o No."),
        ("Ciudad", "Ciudad del cliente."),
        ("Municipio principal", "Municipio principal de domicilio."),
        ("Teléfono del cliente", "Teléfono fijo o celular del cliente."),
        ("Correo del cliente", "Correo electrónico del cliente."),
        ("Tipo de documento", "NIT, CC, CE, TI, Pasaporte o RUT. Si se deja vacío, se asigna NIT."),
        (
            "Correo del contador (asignado en Contanika)",
            "Correo de un usuario en Contanika. Vacío o no encontrado = se asigna a quien importa.",
        ),
        ("Tipos de servicio", "Varios separados por punto y coma. Ej: Declaración de renta; IVA"),
        ("Municipios adicionales", "Varios separados por punto y coma (multi-ICA)."),
    ]
    for fila, (titulo, descripcion) in enumerate(textos, start=1):
        celda = instrucciones.cell(row=fila, column=1, value=titulo)
        if fila == 1:
            celda.font = Font(bold=True)
        instrucciones.cell(row=fila, column=2, value=descripcion).alignment = Alignment(wrap_text=True)
    instrucciones.cell(row=1, column=1).font = Font(bold=True)

    archivo = BytesIO()
    libro.save(archivo)
    archivo.seek(0)
    return archivo


def _normalizar_encabezados(fila_encabezados):
    """Relaciona el título de cada columna con su clave interna."""
    titulos = {titulo.strip().lower(): clave for clave, titulo, _ancho in COLUMNAS}
    mapa = {}
    for indice, valor in enumerate(fila_encabezados):
        if valor is None:
            continue
        clave = titulos.get(str(valor).strip().lower())
        if clave:
            mapa[clave] = indice
    return mapa


def _dividir_lista(valor):
    """Convierte una celda con valores separados por ; o , en una lista limpia."""
    if valor is None:
        return []
    texto = str(valor).replace(",", ";") if ";" not in str(valor) else str(valor)
    return [parte.strip() for parte in texto.split(";") if parte.strip()]


def _leer_fila(valores, mapa, numero_fila):
    """Valida los datos de una fila y devuelve un diccionario listo para guardar."""

    def dato(clave):
        indice = mapa.get(clave)
        if indice is None or indice >= len(valores):
            return ""
        valor = valores[indice]
        return "" if valor is None else str(valor).strip()

    nombre = dato("nombre")
    nit = dato("nit")
    tipo_documento = dato("tipo_documento").upper()
    tipo_persona = dato("tipo_persona").capitalize()

    if not nombre:
        raise ValueError("Falta el nombre o razón social.")
    if not nit:
        raise ValueError("Falta el NIT o cédula.")
    if tipo_documento and tipo_documento not in dict(Cliente.OPCIONES_TIPO_DOCUMENTO):
        raise ValueError("El tipo de documento debe ser NIT, CC, CE, TI, Pasaporte o RUT.")
    if tipo_persona not in (Cliente.NATURAL, Cliente.JURIDICA):
        raise ValueError("El tipo de persona debe ser Natural o Jurídica.")

    responsable_iva = dato("responsable_iva").lower()
    if responsable_iva not in VALORES_VERDADEROS | VALORES_FALSOS:
        raise ValueError("El campo Responsable de IVA debe ser Sí o No.")

    correo_responsable = dato("responsable_correo")
    responsable = None
    if correo_responsable:
        responsable = Usuario.objects.filter(correo__iexact=correo_responsable).first()

    return {
        "fila": numero_fila,
        "nombre": nombre,
        "nit": nit,
        "tipo_documento": tipo_documento or Cliente.NIT,
        "tipo_persona": tipo_persona,
        "regimen": dato("regimen"),
        "responsable_iva": responsable_iva in VALORES_VERDADEROS,
        "ciudad": dato("ciudad"),
        "municipio_principal": dato("municipio_principal"),
        "telefono": dato("telefono"),
        "correo_cliente": dato("correo_cliente"),
        "responsable": responsable,
        "tipos_servicio": _dividir_lista(dato("tipos_servicio")),
        "municipios_adicionales": _dividir_lista(dato("municipios_adicionales")),
    }


def importar_clientes(archivo, usuario):
    """
    Valida e importa los clientes del archivo Excel.

    Devuelve un diccionario con el total de filas leídas, cuántas se
    importaron correctamente y la lista de errores (fila y motivo). Las filas
    válidas se guardan aunque otras tengan error.
    """
    libro = load_workbook(archivo, data_only=True)
    hoja = libro.active
    filas = list(hoja.iter_rows(values_only=True))
    if not filas:
        raise ValueError("El archivo está vacío.")

    mapa = _normalizar_encabezados(filas[0])
    if "nombre" not in mapa or "nit" not in mapa or "tipo_persona" not in mapa:
        raise ValueError("El archivo no corresponde a la plantilla de clientes.")

    resultado = {"total": 0, "importados": 0, "errores": []}
    validos = []
    nits_vistos = set()

    for numero_fila, valores in enumerate(filas[1:], start=2):
        if not any(valor is not None and str(valor).strip() for valor in valores):
            continue
        resultado["total"] += 1
        try:
            datos = _leer_fila(valores, mapa, numero_fila)
            if datos["nit"].lower() in nits_vistos:
                raise ValueError("El NIT o cédula está repetido dentro del archivo.")
            if Cliente.objects.filter(nit_cedula__iexact=datos["nit"]).exists():
                raise ValueError("Ya existe un cliente con ese NIT o cédula en el sistema.")
            nits_vistos.add(datos["nit"].lower())
            validos.append(datos)
        except ValueError as error:
            resultado["errores"].append({"fila": numero_fila, "mensaje": str(error)})

    for datos in validos:
        try:
            with transaction.atomic():
                _crear_cliente(datos, usuario)
            resultado["importados"] += 1
        except Exception as error:  # noqa: BLE001 - se informa el motivo al usuario
            resultado["errores"].append({"fila": datos["fila"], "mensaje": str(error)})

    return resultado


def _crear_cliente(datos, usuario):
    """Guarda un cliente validado con sus servicios y municipios adicionales."""
    cliente = Cliente.objects.create(
        nombre_razon_social=datos["nombre"],
        nit_cedula=datos["nit"],
        tipo_documento=datos["tipo_documento"],
        tipo_persona=datos["tipo_persona"],
        regimen_tributario=datos["regimen"],
        responsable_iva=datos["responsable_iva"],
        ciudad=datos["ciudad"],
        municipio_principal=datos["municipio_principal"],
        informacion_contacto=datos["telefono"],
        correo_cliente=datos["correo_cliente"],
        responsable=datos["responsable"] or usuario,
    )
    for nombre_servicio in datos["tipos_servicio"]:
        tipo, _ = TipoServicio.objects.get_or_create(nombre_servicio=nombre_servicio)
        ClienteServicio.objects.get_or_create(cliente=cliente, tipo_servicio=tipo)
    for municipio in datos["municipios_adicionales"]:
        ClienteMunicipio.objects.get_or_create(cliente=cliente, municipio=municipio)
    return cliente
