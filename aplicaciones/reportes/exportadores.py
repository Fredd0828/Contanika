"""
Exportación de los reportes del módulo 10 a Excel (openpyxl) y PDF (ReportLab).

Ambos formatos parten del mismo `ResultadoReporte`: encabezado con el
despacho, el reporte, la fecha, quién lo generó y los filtros; el resumen;
cada sección como tabla (filas resaltadas en rojo) y el gráfico de barras
cuando el reporte lo tiene.
"""

from datetime import date
from decimal import Decimal
from io import BytesIO

from django.utils import timezone
from django.utils.html import escape
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.barcharts import HorizontalBarChart, VerticalBarChart
from reportlab.graphics.charts.legends import Legend
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import CondPageBreak, KeepTogether, Paragraph, Table, TableStyle

from aplicaciones.comunes import paleta
from aplicaciones.comunes.pdf import (
    ancho_util,
    color,
    construir_pdf,
    encabezado_documento,
    estilo_tabla,
    estilos_documento,
    separador,
)
from aplicaciones.finanzas.templatetags.finanzas_tags import formato_pesos
from aplicaciones.seguridad.models import ConfiguracionDespacho

from .resultados import FECHA, NUMERO, PESOS, PORCENTAJE

# Colores de la paleta de la plataforma (ver `comunes/paleta.py`).
COLOR_ENCABEZADO = paleta.sin_numeral(paleta.TEAL)
COLOR_RESALTADO = "FDE2E1"
COLORES_SERIES = paleta.SERIES_GRAFICO


def formatear(valor, tipo):
    """Texto de una celda para la pantalla y el PDF."""
    if valor is None or valor == "":
        return "—" if tipo in (PESOS, NUMERO, PORCENTAJE, FECHA) else ""
    if tipo == PESOS:
        return formato_pesos(valor)
    if tipo == FECHA and isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    if tipo == PORCENTAJE:
        return f"{valor:+.1f} %".replace(".", ",")
    return str(valor)


def _encabezado(definicion, usuario, filtros_texto):
    despacho = ConfiguracionDespacho.objects.first()
    return {
        "despacho": despacho.nombre_despacho if despacho else "",
        "titulo": f"Reporte de {definicion.nombre.lower()}",
        "generado": timezone.localtime().strftime("%d/%m/%Y %H:%M"),
        "usuario": usuario.nombre,
        "filtros": filtros_texto or "Sin filtros",
        "objeto_despacho": despacho,
    }


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def _valor_excel(valor, tipo):
    if valor is None:
        return None
    if tipo == PESOS:
        return float(valor)
    if tipo == PORCENTAJE:
        return valor / 100
    if isinstance(valor, Decimal):
        return float(valor)
    return valor


def exportar_excel(definicion, resultado, usuario, filtros_texto):
    """Libro de Excel con una hoja por sección y, si aplica, el gráfico nativo de Excel."""
    datos = _encabezado(definicion, usuario, filtros_texto)
    libro = Workbook()
    negrita = Font(bold=True)
    blanco = Font(bold=True, color="FFFFFF")
    relleno_encabezado = PatternFill("solid", fgColor=COLOR_ENCABEZADO)
    relleno_resaltado = PatternFill("solid", fgColor=COLOR_RESALTADO)
    texto_rojo = Font(color="B42318")

    for numero, seccion in enumerate(resultado.secciones):
        hoja = libro.active if numero == 0 else libro.create_sheet()
        hoja.title = seccion.titulo[:31].replace("/", "-").replace(":", "")
        hoja.append([datos["despacho"]])
        hoja.append([datos["titulo"]])
        hoja["A1"].font = Font(bold=True, size=13)
        hoja["A2"].font = Font(bold=True, size=12)
        hoja.append([f"Generado el {datos['generado']} por {datos['usuario']}"])
        hoja.append([f"Filtros: {datos['filtros']}"])
        if numero == 0:
            for etiqueta, valor in resultado.resumen:
                hoja.append([etiqueta, valor])
                hoja.cell(row=hoja.max_row, column=1).font = negrita
        hoja.append([])
        hoja.append([seccion.titulo])
        hoja.cell(row=hoja.max_row, column=1).font = negrita
        hoja.append([columna.titulo for columna in seccion.columnas])
        fila_titulos = hoja.max_row
        for celda in hoja[fila_titulos]:
            celda.font = blanco
            celda.fill = relleno_encabezado
            celda.alignment = Alignment(wrap_text=True, vertical="center")
        for indice, fila in enumerate(seccion.filas):
            hoja.append([_valor_excel(v, c.tipo) for v, c in zip(fila, seccion.columnas, strict=True)])
            for posicion, columna in enumerate(seccion.columnas, start=1):
                celda = hoja.cell(row=hoja.max_row, column=posicion)
                if columna.tipo == PESOS:
                    celda.number_format = '"$" #,##0'
                elif columna.tipo == FECHA:
                    celda.number_format = "DD/MM/YYYY"
                elif columna.tipo == PORCENTAJE:
                    celda.number_format = "0.0%"
                if indice in seccion.resaltadas:
                    celda.fill = relleno_resaltado
                    celda.font = texto_rojo
        if seccion.nota_resaltado and seccion.resaltadas:
            hoja.append([seccion.nota_resaltado])
        for posicion, columna in enumerate(seccion.columnas, start=1):
            hoja.column_dimensions[get_column_letter(posicion)].width = columna.ancho + 2
        hoja.freeze_panes = hoja.cell(row=fila_titulos + 1, column=1)

        if numero == 0 and resultado.grafico and seccion.filas:
            _grafico_excel(libro, resultado.grafico)

    for nota in resultado.notas:
        libro.worksheets[0].append([nota])

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _grafico_excel(libro, grafico):
    """Hoja «Gráfico» con los datos y un gráfico de barras nativo de Excel."""
    hoja = libro.create_sheet("Gráfico")
    hoja.append(["Categoría"] + [nombre for nombre, _ in grafico.series])
    for posicion, etiqueta in enumerate(grafico.etiquetas):
        hoja.append([etiqueta] + [valores[posicion] for _, valores in grafico.series])
    barras = BarChart()
    barras.type = "bar" if grafico.horizontal else "col"
    barras.title = grafico.titulo
    barras.y_axis.title = "Pesos"
    barras.height, barras.width = 10, 22
    datos = Reference(
        hoja, min_col=2, max_col=1 + len(grafico.series), min_row=1, max_row=1 + len(grafico.etiquetas)
    )
    categorias = Reference(hoja, min_col=1, min_row=2, max_row=1 + len(grafico.etiquetas))
    barras.add_data(datos, titles_from_data=True)
    barras.set_categories(categorias)
    # Sin esto Excel pinta las barras con el azul de su tema por defecto.
    for posicion, serie in enumerate(barras.series):
        color = paleta.sin_numeral(COLORES_SERIES[posicion % len(COLORES_SERIES)])
        serie.graphicalProperties.solidFill = color
        serie.graphicalProperties.line.solidFill = color
    hoja.add_chart(barras, f"{get_column_letter(len(grafico.series) + 3)}2")
    hoja.column_dimensions["A"].width = 30


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


TIPOS_A_LA_DERECHA = (PESOS, NUMERO, PORCENTAJE)
MARGEN_REPORTE = 1.5 * cm


def _dibujo_grafico(grafico, ancho):
    """Gráfico de barras de ReportLab, con la letra y los colores del documento."""
    alto = 7 * cm if not grafico.horizontal else max(4.5 * cm, 0.7 * cm * len(grafico.etiquetas) + 1.6 * cm)
    varias_series = len(grafico.series) > 1
    dibujo = Drawing(ancho, alto)
    clase = HorizontalBarChart if grafico.horizontal else VerticalBarChart
    barras = clase()
    margen_izq = 5.6 * cm if grafico.horizontal else 2.4 * cm
    margen_der = 3.8 * cm if varias_series else 0.8 * cm
    barras.x, barras.y = margen_izq, 1.1 * cm
    barras.width, barras.height = ancho - margen_izq - margen_der, alto - 1.5 * cm
    barras.data = [valores for _, valores in grafico.series]
    barras.barSpacing = 1
    barras.groupSpacing = 6
    tenue = color(paleta.TEXTO_TENUE)
    linea = color(paleta.BORDE)
    eje_categorias = barras.categoryAxis
    eje_categorias.categoryNames = grafico.etiquetas
    eje_categorias.labels.fontName = "Helvetica"
    eje_categorias.labels.fontSize = 7.5
    eje_categorias.labels.fillColor = tenue
    eje_categorias.strokeColor = linea
    eje_categorias.visibleTicks = False
    if not grafico.horizontal:
        eje_categorias.labels.angle = 30
        eje_categorias.labels.dy = -6
    eje_valores = barras.valueAxis
    eje_valores.valueMin = 0
    eje_valores.labels.fontName = "Helvetica"
    eje_valores.labels.fontSize = 7.5
    eje_valores.labels.fillColor = tenue
    eje_valores.strokeColor = linea
    eje_valores.visibleTicks = False
    eje_valores.visibleGrid = True
    eje_valores.gridStrokeColor = color(paleta.BORDE_SUAVE)
    eje_valores.gridStrokeWidth = 0.5
    eje_valores.labelTextFormat = lambda valor: formato_pesos(valor)
    for posicion, _ in enumerate(grafico.series):
        barras.bars[posicion].fillColor = colors.HexColor(COLORES_SERIES[posicion % len(COLORES_SERIES)])
        barras.bars[posicion].strokeColor = None
    dibujo.add(barras)
    if varias_series:
        leyenda = Legend()
        leyenda.x, leyenda.y = ancho - 3.4 * cm, alto - 0.6 * cm
        leyenda.fontName = "Helvetica"
        leyenda.fontSize = 7.5
        leyenda.fillColor = tenue
        leyenda.dx = leyenda.dy = 7
        leyenda.deltay = 11
        leyenda.strokeWidth = 0
        leyenda.alignment = "right"
        leyenda.colorNamePairs = [
            (colors.HexColor(COLORES_SERIES[i % len(COLORES_SERIES)]), nombre)
            for i, (nombre, _) in enumerate(grafico.series)
        ]
        dibujo.add(leyenda)
    return dibujo


def _tarjetas_resumen(resumen, ancho, estilos):
    """Las cifras del resumen en una fila de tarjetas: etiqueta pequeña y valor grande."""
    celdas = [
        [
            Paragraph(escape(str(etiqueta)).upper(), estilos["etiqueta"]),
            Paragraph(escape(str(valor)), estilos["cifra"]),
        ]
        for etiqueta, valor in resumen
    ]
    tabla = Table([celdas], colWidths=[ancho / len(celdas)] * len(celdas))
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), color(paleta.SUPERFICIE)),
                ("BOX", (0, 0), (-1, -1), 0.6, color(paleta.BORDE)),
                ("LINEAFTER", (0, 0), (-2, -1), 0.4, color(paleta.BORDE_SUAVE)),
                ("LINEABOVE", (0, 0), (-1, 0), 1.6, color(paleta.TEAL)),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 9),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
                ("LEFTPADDING", (0, 0), (-1, -1), 11),
                ("RIGHTPADDING", (0, 0), (-1, -1), 11),
            ]
        )
    )
    return tabla


def _tabla_seccion(seccion, ancho, estilos):
    """Tabla de una sección: cifras a la derecha, renglón alterno y filas en mora en rojo."""
    total_ancho = sum(columna.ancho for columna in seccion.columnas)
    anchos = [ancho * columna.ancho / total_ancho for columna in seccion.columnas]
    derecha = [columna.tipo in TIPOS_A_LA_DERECHA for columna in seccion.columnas]
    filas = [
        [
            Paragraph(
                escape(columna.titulo), estilos["celda_titulo_derecha" if a_la_derecha else "celda_titulo"]
            )
            for columna, a_la_derecha in zip(seccion.columnas, derecha, strict=True)
        ]
    ]
    roja = ParagraphStyle("celda_roja", parent=estilos["celda"], textColor=color(paleta.MARRON))
    roja_derecha = ParagraphStyle(
        "celda_roja_derecha", parent=estilos["celda_derecha"], textColor=color(paleta.MARRON)
    )
    for indice, fila in enumerate(seccion.filas):
        resaltada = indice in seccion.resaltadas
        celdas = []
        for valor, columna, a_la_derecha in zip(fila, seccion.columnas, derecha, strict=True):
            if resaltada:
                estilo = roja_derecha if a_la_derecha else roja
            else:
                estilo = estilos["celda_derecha" if a_la_derecha else "celda"]
            celdas.append(Paragraph(escape(formatear(valor, columna.tipo)), estilo))
        filas.append(celdas)
    if not seccion.filas:
        filas.append(
            [Paragraph("Sin registros para los filtros elegidos.", estilos["celda"])]
            + [""] * (len(anchos) - 1)
        )
    tabla = Table(filas, colWidths=anchos, repeatRows=1)
    estilo = estilo_tabla(resaltadas=seccion.resaltadas, total_filas=len(filas) - 1)
    if not seccion.filas:
        estilo.add("SPAN", (0, 1), (-1, 1))
    tabla.setStyle(estilo)
    return tabla


def exportar_pdf(definicion, resultado, usuario, filtros_texto):
    """
    PDF horizontal del reporte, con la hoja común de la plataforma: encabezado
    del despacho y título, filtros aplicados, resumen en tarjetas, gráfico y
    una tabla por sección. Cada título de sección se mantiene junto a su tabla.
    """
    datos = _encabezado(definicion, usuario, filtros_texto)
    tamano = landscape(letter)
    ancho = ancho_util(tamano, MARGEN_REPORTE)
    estilos = estilos_documento()

    elementos = encabezado_documento(
        datos["objeto_despacho"],
        ancho,
        escape(datos["titulo"]),
        [f"Generado el {datos['generado']}", f"Por {escape(datos['usuario'])}"],
        estilos,
    )
    elementos.append(Paragraph(f"<b>Filtros aplicados:</b> {escape(datos['filtros'])}", estilos["nota"]))
    elementos.append(separador(0.35))
    if resultado.resumen:
        elementos += [_tarjetas_resumen(resultado.resumen, ancho, estilos), separador(0.45)]
    if resultado.grafico and resultado.grafico.etiquetas:
        elementos.append(
            KeepTogether(
                [
                    Paragraph(escape(resultado.grafico.titulo), estilos["subtitulo"]),
                    _dibujo_grafico(resultado.grafico, ancho),
                    separador(0.3),
                ]
            )
        )

    for seccion in resultado.secciones:
        # Si no cabe el título con al menos unas filas, la sección empieza en hoja nueva.
        elementos.append(CondPageBreak(3.2 * cm))
        registros = len(seccion.filas)
        cuenta = f"{registros} registro{'s' if registros != 1 else ''}"
        elementos.append(
            Paragraph(
                f'{escape(seccion.titulo)} <font size="8" color="{paleta.TEXTO_APAGADO}">· {cuenta}</font>',
                estilos["subtitulo"],
            )
        )
        elementos.append(_tabla_seccion(seccion, ancho, estilos))
        if seccion.nota_resaltado and seccion.resaltadas:
            elementos += [separador(0.12), Paragraph(escape(seccion.nota_resaltado), estilos["nota_roja"])]
        elementos.append(separador(0.45))
    for nota in resultado.notas:
        elementos.append(Paragraph(escape(nota), estilos["nota"]))

    return construir_pdf(
        elementos,
        tamano=tamano,
        titulo=datos["titulo"],
        despacho=datos["objeto_despacho"],
        margen_lateral=MARGEN_REPORTE,
    )
