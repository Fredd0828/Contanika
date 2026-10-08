"""Pruebas de la importación masiva de la tabla DIAN y del comando de carga."""

from datetime import date
from io import BytesIO

import pytest
from django.core.management import call_command
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from aplicaciones.calendario.importacion import (
    _parse_fecha,
    generar_plantilla_dian,
    importar_dian,
)
from aplicaciones.calendario.models import TablaReferenciaDian

pytestmark = pytest.mark.django_db


def _excel_con_filas(filas):
    """Construye un Excel en memoria con las filas dadas."""
    libro = Workbook()
    hoja = libro.active
    hoja.append(["tipo_obligacion", "grupo_nit", "fecha_vencimiento", "periodo_gravable"])
    for fila in filas:
        hoja.append(fila)
    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return buffer


# ---------------------------------------------------------------------------
# Plantilla descargable
# ---------------------------------------------------------------------------


def test_plantilla_dian_es_excel_valido():
    """La plantilla tiene los encabezados y la hoja de instrucciones."""
    buffer = generar_plantilla_dian()
    libro = load_workbook(buffer)
    assert "calendario_dian" in libro.sheetnames
    assert "instrucciones" in libro.sheetnames
    hoja = libro["calendario_dian"]
    assert hoja["A1"].value == "tipo_obligacion"
    assert hoja["D1"].value == "periodo_gravable"


# ---------------------------------------------------------------------------
# _parse_fecha
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "valor, esperado",
    [
        ("2026-03-10", date(2026, 3, 10)),
        ("10/03/2026", date(2026, 3, 10)),
        ("10-03-2026", date(2026, 3, 10)),
        (date(2026, 3, 10), date(2026, 3, 10)),
        ("texto", None),
        (None, None),
    ],
)
def test_parse_fecha_formatos(valor, esperado):
    """La fecha se acepta en varios formatos y se rechaza la inválida."""
    assert _parse_fecha(valor) == esperado


# ---------------------------------------------------------------------------
# importar_dian
# ---------------------------------------------------------------------------


def test_importar_dian_carga_filas_validas():
    """Las filas válidas se cargan en la tabla de referencia."""
    archivo = _excel_con_filas(
        [
            ["IVA bimestral", "1", "2026-03-10", "2026-01"],
            ["IVA bimestral", "2", "2026-03-11", "2026-01"],
        ]
    )
    resultado = importar_dian(archivo)

    assert resultado["leidas"] == 2
    assert resultado["cargadas"] == 2
    assert resultado["errores"] == []
    assert TablaReferenciaDian.objects.count() == 2


def test_importar_dian_actualiza_duplicados():
    """Una fila ya existente actualiza la fecha en lugar de duplicar."""
    TablaReferenciaDian.objects.create(
        tipo_obligacion="IVA bimestral",
        grupo_nit="1",
        fecha_vencimiento=date(2026, 3, 10),
        periodo_gravable="2026-01",
    )
    archivo = _excel_con_filas([["IVA bimestral", "1", "2026-03-15", "2026-01"]])
    resultado = importar_dian(archivo)

    assert resultado["cargadas"] == 0
    assert resultado["omitidas"] == 1
    referencia = TablaReferenciaDian.objects.get()
    assert referencia.fecha_vencimiento == date(2026, 3, 15)


def test_importar_dian_reporta_filas_incompletas():
    """Las filas con datos faltantes se reportan con su número."""
    archivo = _excel_con_filas(
        [
            ["IVA bimestral", "1", "2026-03-10", "2026-01"],
            ["IVA bimestral", "", "2026-03-11", "2026-01"],
            ["", "3", "fecha-mala", ""],
        ]
    )
    resultado = importar_dian(archivo)

    assert resultado["cargadas"] == 1
    assert len(resultado["errores"]) == 2
    assert "Fila 3" in resultado["errores"][0]


def test_importar_dian_omite_filas_vacias():
    """Las filas completamente vacías no cuentan como leídas."""
    archivo = _excel_con_filas(
        [
            ["IVA bimestral", "1", "2026-03-10", "2026-01"],
            [None, None, None, None],
        ]
    )
    resultado = importar_dian(archivo)
    assert resultado["leidas"] == 1


def test_importar_dian_archivo_corrupto():
    """Un archivo que no es Excel devuelve un error controlado."""
    with pytest.raises(ValueError, match="No se pudo leer"):
        importar_dian(BytesIO(b"esto no es un excel"))


# ---------------------------------------------------------------------------
# Vista de importación
# ---------------------------------------------------------------------------


def test_vista_importar_dian_con_excel(client, usuario_administrador):
    """El Administrador importa el Excel desde la vista."""
    from django.core.files.uploadedfile import SimpleUploadedFile

    client.force_login(usuario_administrador)
    contenido = _excel_con_filas([["IVA bimestral", "1", "2026-03-10", "2026-01"]])
    archivo = SimpleUploadedFile(
        "calendario.xlsx",
        contenido.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    respuesta = client.post(
        reverse("calendario:importar_dian"),
        {"archivo": archivo},
        format="multipart",
    )
    assert respuesta.status_code == 302
    assert TablaReferenciaDian.objects.count() == 1


def test_vista_importar_dian_solo_administrador(client, usuario_senior):
    """El Contador Senior no puede importar la tabla DIAN."""
    from django.core.files.uploadedfile import SimpleUploadedFile

    client.force_login(usuario_senior)
    contenido = _excel_con_filas([["IVA bimestral", "1", "2026-03-10", "2026-01"]])
    archivo = SimpleUploadedFile("calendario.xlsx", contenido.getvalue())
    respuesta = client.post(
        reverse("calendario:importar_dian"),
        {"archivo": archivo},
        format="multipart",
    )
    assert respuesta.status_code == 403


# ---------------------------------------------------------------------------
# Management command de carga del calendario oficial
# ---------------------------------------------------------------------------


def test_comando_cargar_calendario_dian_2026():
    """El comando carga los 440 registros oficiales y es idempotente."""
    call_command("cargar_calendario_dian_2026")
    total = TablaReferenciaDian.objects.count()
    assert total == 440

    call_command("cargar_calendario_dian_2026")
    assert TablaReferenciaDian.objects.count() == 440


def test_comando_carga_fechas_correctas():
    """Verifica dos fechas oficiales tomadas del calendario DIAN 2026."""
    call_command("cargar_calendario_dian_2026")

    iva = TablaReferenciaDian.objects.get(
        tipo_obligacion="IVA bimestral", grupo_nit="6", periodo_gravable="2026-01"
    )
    assert iva.fecha_vencimiento == date(2026, 3, 17)

    renta = TablaReferenciaDian.objects.get(
        tipo_obligacion="Renta personas naturales", grupo_nit="56", periodo_gravable="2026"
    )
    assert renta.fecha_vencimiento == date(2026, 9, 21)
