"""Pruebas de los seis reportes del módulo 10 y de su exportación a Excel y PDF."""

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from aplicaciones.reportes import exportadores
from aplicaciones.reportes.definiciones import POR_SLUG, REPORTES
from aplicaciones.reportes.formularios import FiltroVencimientosForm

pytestmark = pytest.mark.django_db

HOY = date(2026, 10, 1)


def _generar(slug, usuario, **filtros):
    return POR_SLUG[slug].generar(usuario, filtros, hoy=HOY)


def test_hay_seis_reportes():
    assert len(REPORTES) == 6


def test_clientes_administrador_ve_todos_y_junior_los_suyos(
    usuario_administrador, usuario_junior, datos_reportes
):
    assert _generar("clientes", usuario_administrador).total_filas == 2
    resultado = _generar("clientes", usuario_junior)
    assert [fila[0] for fila in resultado.secciones[0].filas] == ["Papelería El Lápiz"]


def test_clientes_filtra_por_regimen_y_estado(usuario_administrador, datos_reportes, cliente):
    cliente.estado = "Inactivo"
    cliente.save()
    resultado = _generar("clientes", usuario_administrador, estado="Inactivo")
    assert [fila[0] for fila in resultado.secciones[0].filas] == [cliente.nombre_razon_social]


def test_vencimientos_marca_en_rojo_los_vencidos_sin_presentar(usuario_administrador, datos_reportes):
    seccion = _generar("vencimientos", usuario_administrador).secciones[0]
    assert [fila[1] for fila in seccion.filas] == ["IVA bimestral", "Retención en la fuente"]
    assert seccion.resaltadas == {0}


def test_vencimientos_del_senior_solo_de_sus_clientes(usuario_senior, datos_reportes, cliente):
    assert _generar("vencimientos", usuario_senior).total_filas == 0
    cliente.responsable = usuario_senior
    cliente.save()
    assert _generar("vencimientos", usuario_senior).total_filas == 1


def test_cartera_con_honorario_deuda_y_abonos(usuario_administrador, datos_reportes):
    resultado = _generar("cartera", usuario_administrador)
    cartera, abonos = resultado.secciones
    fila = next(f for f in cartera.filas if f[0] == "Comercializadora La Rebaja SAS")
    assert fila[2:5] == [Decimal("1000000"), Decimal("400000"), Decimal("600000")]
    assert fila[7] == "Crítica"  # venció el 10 de agosto: más de 30 días de mora
    assert len(abonos.filas) == 1 and abonos.filas[0][4] == Decimal("400000")
    assert resultado.grafico.horizontal is True


def test_cuentas_de_cobro_filtra_por_estado_de_pago_y_marca_vencidas(usuario_administrador, datos_reportes):
    # Venció el 10 de agosto con saldo: el disparador la dejó "Vencida" al registrar el abono.
    seccion = _generar("cuentas-cobro", usuario_administrador, estado_pago="Vencida").secciones[0]
    assert [fila[0] for fila in seccion.filas] == [datos_reportes["cuenta"].pk]
    assert seccion.resaltadas == {0}


def test_tareas_incluye_historial_y_marca_vencidas(usuario_administrador, datos_reportes):
    seccion = _generar("tareas", usuario_administrador).secciones[0]
    vencida = next(i for i, f in enumerate(seccion.filas) if f[1] == "Conciliar bancos")
    assert vencida in seccion.resaltadas


def test_ingresos_por_mes_comparativo_y_top3(usuario_administrador, datos_reportes):
    resultado = _generar("ingresos", usuario_administrador, desde=date(2026, 8, 1), hasta=date(2026, 10, 31))
    meses = resultado.secciones[0].filas
    assert [fila[0] for fila in meses] == ["Ago 2026", "Sep 2026", "Oct 2026"]
    septiembre = meses[1]
    assert septiembre[2] == Decimal("400000")  # cobrado: abono del 5 de septiembre
    facturado_total = sum(fila[1] for fila in meses)
    assert facturado_total == Decimal("1500000")
    assert resultado.secciones[1].titulo.startswith("Comparativo")
    top = resultado.secciones[2].filas
    assert top[0][0] == "Comercializadora La Rebaja SAS"
    assert [nombre for nombre, _ in resultado.grafico.series] == [
        "Facturado",
        "Cobrado",
        "Pendiente",
        "En mora",
    ]


def test_ingresos_es_consolidado_aunque_lo_pida_un_junior(usuario_junior, datos_reportes):
    resultado = _generar("ingresos", usuario_junior, desde=date(2026, 8, 1), hasta=date(2026, 10, 31))
    assert sum(fila[1] for fila in resultado.secciones[0].filas) == Decimal("1500000")


def test_filtro_rechaza_rango_invertido(usuario_administrador):
    formulario = FiltroVencimientosForm(
        {"desde": "2026-12-01", "hasta": "2026-01-01"}, usuario=usuario_administrador
    )
    assert not formulario.is_valid()


def test_resumen_de_filtros_legible(usuario_administrador, cliente):
    formulario = FiltroVencimientosForm(
        {"desde": "2026-01-01", "cliente": cliente.pk, "ambito": "Nacional"}, usuario=usuario_administrador
    )
    texto = formulario.resumen_filtros()
    assert (
        "Desde: 01/01/2026" in texto and "Cliente: Comercializadora" in texto and "Ámbito: Nacional" in texto
    )


# ---------------------------------------------------------------------------
# Exportación
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("slug", list(POR_SLUG))
def test_exporta_excel_y_pdf(slug, usuario_administrador, datos_reportes):
    definicion = POR_SLUG[slug]
    resultado = definicion.generar(usuario_administrador, {}, hoy=HOY)
    excel = exportadores.exportar_excel(definicion, resultado, usuario_administrador, "Sin filtros")
    libro = load_workbook(BytesIO(excel))
    assert libro.worksheets[0]["A2"].value == f"Reporte de {definicion.nombre.lower()}"
    if resultado.grafico:
        assert "Gráfico" in libro.sheetnames
        assert libro["Gráfico"]._charts
    pdf = exportadores.exportar_pdf(definicion, resultado, usuario_administrador, "Sin filtros")
    assert pdf.startswith(b"%PDF")


def test_excel_usa_numeros_y_marca_filas_en_rojo(usuario_administrador, datos_reportes):
    definicion = POR_SLUG["vencimientos"]
    resultado = definicion.generar(usuario_administrador, {}, hoy=HOY)
    hoja = load_workbook(
        BytesIO(exportadores.exportar_excel(definicion, resultado, usuario_administrador, ""))
    ).active
    fila_iva = next(fila for fila in hoja.iter_rows() if fila[1].value == "IVA bimestral")
    assert fila_iva[0].fill.fgColor.rgb.endswith(exportadores.COLOR_RESALTADO)
    assert fila_iva[5].number_format == "DD/MM/YYYY"


def test_formatear_valores():
    assert exportadores.formatear(Decimal("1200000"), "pesos") == "$ 1.200.000"
    assert exportadores.formatear(date(2026, 3, 1), "fecha") == "01/03/2026"
    assert exportadores.formatear(12.5, "porcentaje") == "+12,5 %"
    assert exportadores.formatear(None, "pesos") == "—"
